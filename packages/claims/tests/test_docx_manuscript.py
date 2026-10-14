"""A Word manuscript is checked as a text manuscript is.

Each test builds its `.docx` with `python-docx` in the test's own directory, so no binary is
committed and the fixture is exactly what the assertions describe. The text is `citations`'
`docx-text` version 2, whose own tests cover each kind of content it reads.
"""

from __future__ import annotations

import json
from importlib.metadata import version

import docx
from citations.verify import extract_uncached
from claims.cli import main
from claims.models import (
    ArtifactRef,
    Claim,
    Digest,
    Manifest,
    MetricEvidence,
    Outcome,
    ProseLocator,
    QuoteEvidence,
    Reason,
    Standing,
    ValueEvidence,
)
from claims.numbers import numbers
from claims.verify import verify
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SENTENCE = "Thirty-five participants responded to the survey."


def revision(kind: str, words: str) -> OxmlElement:
    """A tracked insertion or deletion of `words`, as Word writes one."""
    change = OxmlElement(f"w:{kind}")
    change.set(qn("w:id"), "1")
    change.set(qn("w:author"), "a")
    run, text = OxmlElement("w:r"), OxmlElement("w:t" if kind == "ins" else "w:delText")
    text.text = words
    text.set(qn("xml:space"), "preserve")
    run.append(text)
    change.append(run)
    return change


def manuscript(path, *, tracked=False):
    """A `.docx` of two paragraphs, a table of two rows, one cell of which is merged, and a
    header. `tracked` inserts ` on average` and deletes ` in 2020` as tracked changes."""
    document = docx.Document()
    document.add_paragraph(SENTENCE)
    paragraph = document.add_paragraph("Responders were older by 4.2 years")
    if tracked:
        paragraph._p.append(revision("ins", " on average"))
        paragraph._p.append(revision("del", " in 2020"))
    paragraph.add_run(".")
    table = document.add_table(rows=2, cols=3)
    for column, heading in enumerate(("Arm", "Cohort", "n")):
        table.cell(0, column).text = heading
    table.cell(1, 0).text = "Treated"
    table.cell(1, 1).merge(table.cell(1, 2)).text = "1,204"
    document.sections[0].header.paragraphs[0].text = "Running head 99"
    document.save(path)
    return path


def project(root, claims, *, tracked=False):
    paper = manuscript(root / "paper.docx", tracked=tracked)
    results = root / "results.json"
    results.write_text(json.dumps({"responders": 35, "treated": 1204}))
    made = Manifest(
        project="p",
        path=root / "repro.yaml",
        artifacts=tuple(
            ArtifactRef(id=p.stem, path=p, digest=Digest.of_file(p)) for p in (paper, results)
        ),
        claims=tuple(claims),
    )
    made.path.write_text(made.model_dump_json())
    return made


def quoted(name, passage, reported="", pointer=""):
    evidence = [QuoteEvidence(artifact="paper", text=passage)]
    if reported:
        evidence.append(
            MetricEvidence(artifact="results", name=name, reported=reported, pointer=pointer)
        )
    return Claim(id=name, text=passage, evidence=tuple(evidence))


def test_the_docx_is_read_in_document_order_and_its_header_after_the_body(tmp_path):
    assert extract_uncached(manuscript(tmp_path / "paper.docx")).splitlines() == [
        SENTENCE,
        "Responders were older by 4.2 years.",
        "Arm | Cohort | n",
        "Treated | 1,204",
        "Running head 99",
    ]


def test_a_passage_in_the_docx_verifies_and_one_it_does_not_hold_is_a_mismatch(tmp_path):
    made = project(
        tmp_path,
        [
            quoted("present", "participants responded to the survey"),
            quoted("row", "Treated | 1,204"),
            quoted("absent", "participants declined the survey"),
            quoted("header", "Running head 99"),
        ],
    )
    outcomes = {d.claim_id: (d.outcome, d.reason) for d in verify(made).decisions}
    assert outcomes == {
        "present": (Outcome.VERIFIED, Reason.PASSAGE_PRESENT),
        "row": (Outcome.VERIFIED, Reason.PASSAGE_PRESENT),
        "absent": (Outcome.MISMATCH, Reason.PASSAGE_ABSENT),
        "header": (Outcome.VERIFIED, Reason.PASSAGE_PRESENT),
    }


def test_a_decision_on_a_docx_names_the_reader_and_its_version_and_hashes_its_text(tmp_path):
    made = project(
        tmp_path,
        [
            quoted("present", "participants responded to the survey"),
            Claim(
                id="older",
                text="t",
                evidence=(
                    ValueEvidence(
                        artifact="paper",
                        name="older",
                        reported="4.2",
                        locator=ProseLocator(before="older by", after="years"),
                    ),
                ),
            ),
            quoted("responders", SENTENCE, "35", "/responders"),
        ],
    )
    quote, prose, _, metric = verify(made).decisions
    text = extract_uncached(tmp_path / "paper.docx")
    # `docx-text` is `citations`' own code, so its version is the one that says which reading.
    assert (quote.tool, quote.tool_version) == ("citations", version("citations"))
    assert quote.extraction_digest == Digest.of_text(text).value
    assert prose.outcome is Outcome.VERIFIED
    assert (prose.tool, prose.tool_version) == ("citations", version("citations"))
    # The JSON beside it is read by the adapters, not by what reads the docx.
    assert metric.tool != "citations"


def test_a_number_in_a_docx_table_is_found_and_bound_by_the_claim_that_quotes_its_row(
    tmp_path, capsys
):
    made = project(tmp_path, [quoted("treated", "Treated | 1,204", "1204", "/treated")])
    (inventory,) = numbers(made).texts
    assert [(n.text, n.standing) for n in inventory.numbers] == [
        ("thirty-five", Standing.UNBOUND),
        ("4.2", Standing.UNBOUND),
        ("1,204", Standing.BOUND),
        ("99", Standing.UNBOUND),
    ]
    assert inventory.numbers[2].claims == ("treated",)

    assert main(["numbers", str(made.path)]) == 0
    assert (
        "4 numbers: 1 bound, 0 set aside by rule, 0 set aside with a reason, 3 unbound"
        in capsys.readouterr().out
    )


def test_a_tracked_insertion_is_read_and_a_tracked_deletion_is_not(tmp_path):
    made = project(
        tmp_path,
        [
            quoted("inserted", "older by 4.2 years on average", "35", "/responders"),
            quoted("deleted", "older by 4.2 years in 2020"),
        ],
        tracked=True,
    )
    inserted, _, deleted = verify(made).decisions
    assert (inserted.outcome, deleted.outcome) == (Outcome.VERIFIED, Outcome.MISMATCH)
    (inventory,) = numbers(made).texts
    assert "2020" not in [n.text for n in inventory.numbers]
