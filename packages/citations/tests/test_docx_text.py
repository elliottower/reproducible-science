"""`docx-text` version 2 reads a `.docx` from its own XML.

Each document is written part by part into a zip in the test's directory, so a fixture holds
exactly the markup a test is about and no binary is committed. A claims file pinned under
version 1 goes on being read by version 1.
"""

from __future__ import annotations

import hashlib
import zipfile

import docx
import pytest
import yaml
from citations import cli, pin
from citations import extractors as E
from citations import verify as V
from citations.models import load_claim_file

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
RELS = "http://schemas.openxmlformats.org/package/2006/relationships"
KINDS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS = (
    f'xmlns:w="{W}" xmlns:r="{R}" '
    'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" '
    'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
    'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
    'xmlns:v="urn:schemas-microsoft-com:vml"'
)
TYPES = (
    '<?xml version="1.0" encoding="UTF-8"?>'
    '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
    '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
    '<Default Extension="xml" ContentType="application/xml"/>'
    '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-'
    'officedocument.wordprocessingml.document.main+xml"/>'
    "</Types>"
)


@pytest.fixture(autouse=True)
def _no_cache():
    V.clear_caches()


def p(*runs: str, mark: str = "") -> str:
    return f"<w:p>{mark}{''.join(runs)}</w:p>"


def r(text: str) -> str:
    return f'<w:r><w:t xml:space="preserve">{text}</w:t></w:r>'


def cell(*paragraphs: str) -> str:
    return f"<w:tc>{''.join(paragraphs)}</w:tc>"


def row(*cells: str, deleted: bool = False) -> str:
    properties = '<w:trPr><w:del w:id="9" w:author="a"/></w:trPr>' if deleted else ""
    return f"<w:tr>{properties}{''.join(cells)}</w:tr>"


def table(*rows: str) -> str:
    """A table with the column grid Word writes, which `python-docx` 1.1 needs to read a row:
    one column per cell of the first row, not counting a table nested in it."""
    columns = rows[0].split("<w:tbl>")[0].count("<w:tc>") if rows else 0
    grid = "<w:tblGrid>" + '<w:gridCol w:w="2000"/>' * columns + "</w:tblGrid>"
    return f"<w:tbl>{grid}{''.join(rows)}</w:tbl>"


def note(kind: str, number: int, text: str) -> str:
    return f'<w:{kind} w:id="{number}">{p(r(text))}</w:{kind}>'


def write(path, body: str, parts: dict[str, str] | None = None, rels: dict[str, str] | None = None):
    """A `.docx` of `body`, with each extra part named in `parts` and related by `rels`, which
    maps a relationship id to `type target`."""
    related = "".join(
        f'<Relationship Id="{rid}" Type="{KINDS}/{kind}" Target="{target}"/>'
        for rid, (kind, target) in ((k, v.split()) for k, v in (rels or {}).items())
    )
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("[Content_Types].xml", TYPES)
        z.writestr(
            "_rels/.rels",
            f'<Relationships xmlns="{RELS}"><Relationship Id="rId1" '
            f'Type="{KINDS}/officeDocument" Target="word/document.xml"/></Relationships>',
        )
        z.writestr("word/document.xml", f"<w:document {NS}><w:body>{body}</w:body></w:document>")
        z.writestr(
            "word/_rels/document.xml.rels",
            f'<Relationships xmlns="{RELS}">{related}</Relationships>',
        )
        for name, xml in (parts or {}).items():
            z.writestr(name, xml)
    return path


def read(path) -> list[str]:
    return E.Declared(E.EXTRACTORS["docx-text"]).read(path).splitlines()


def manuscript(path):
    """Every kind of content version 2 reads, and a comment it does not."""
    footnote_ref = '<w:r><w:footnoteReference w:id="2"/></w:r>'
    endnote_ref = '<w:r><w:endnoteReference w:id="5"/></w:r>'
    textbox = (
        '<w:r><mc:AlternateContent><mc:Choice Requires="wps"><w:drawing><wps:txbx>'
        f"<w:txbxContent>{p(r('A text box beside the figure.'))}</w:txbxContent>"
        "</wps:txbx></w:drawing></mc:Choice><mc:Fallback><w:pict><v:textbox>"
        f"<w:txbxContent>{p(r('A text box beside the figure.'))}</w:txbxContent>"
        "</v:textbox></w:pict></mc:Fallback></mc:AlternateContent></w:r>"
    )
    complex_field = (
        '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
        '<w:r><w:instrText> ZOTERO_ITEM CSL_CITATION {"citationID":"x"} </w:instrText></w:r>'
        '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
        + r("(Smith 2020)")
        + '<w:r><w:fldChar w:fldCharType="end"/></w:r>'
    )
    body = "".join(
        [
            p(r("Thirty-five participants responded"), footnote_ref, r(".")),
            table(
                row(cell(p(r("Arm"))), cell(p(r("n")))),
                row(
                    cell(p(r("Treated"))),
                    cell(p(r("1,204")), table(row(cell(p(r("inner"))), cell(p(r("7")))))),
                ),
                row(cell(p(r("Withdrawn"))), cell(p(r("3"))), deleted=True),
            ),
            p(r("Responders were older"), endnote_ref, r(" by 4.2 years."), textbox),
            '<w:sdt><w:sdtPr><w:alias w:val="Abstract"/></w:sdtPr><w:sdtContent>'
            + p(r("A block inside a content control."))
            + "</w:sdtContent></w:sdt>",
            p(
                r("As reported "),
                complex_field,
                r(" and "),
                "<w:sdt><w:sdtPr/><w:sdtContent>" + r("(Jones 2021)") + "</w:sdtContent></w:sdt>",
                r(", see "),
                f'<w:hyperlink r:id="rId9">{r("the registry")}</w:hyperlink>',
                r(" in table "),
                f'<w:fldSimple w:instr=" SEQ Table ">{r("12")}</w:fldSimple>',
                r("."),
            ),
            p(
                r("The effect "),
                '<w:ins w:id="1" w:author="a">' + r("held in every cohort ") + "</w:ins>",
                '<w:del w:id="2" w:author="a"><w:r><w:delText>vanished </w:delText></w:r></w:del>',
                r("at "),
                "<m:oMath><m:r><m:t>p</m:t></m:r><m:r><m:t>&lt;0.05</m:t></m:r></m:oMath>",
                r("."),
                '<w:commentRangeStart w:id="0"/>',
                '<w:r><w:commentReference w:id="0"/></w:r>',
            ),
            p(
                r("A paragraph whose mark is deleted "),
                mark='<w:pPr><w:rPr><w:del w:id="3" w:author="a"/></w:rPr></w:pPr>',
            ),
            p(r("joins the next.")),
            '<w:p><w:pPr><w:sectPr><w:headerReference w:type="default" r:id="rIdH"/>'
            '<w:footerReference w:type="default" r:id="rIdF"/></w:sectPr></w:pPr></w:p>',
            '<w:sectPr><w:headerReference w:type="default" r:id="rIdH"/>'
            '<w:footerReference w:type="default" r:id="rIdF"/></w:sectPr>',
        ]
    )
    parts = {
        "word/footnotes.xml": f"<w:footnotes {NS}>"
        '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
        + note("footnote", 2, "Participants were paid for their time.")
        + "</w:footnotes>",
        "word/endnotes.xml": f"<w:endnotes {NS}>"
        + note("endnote", 5, "Age at enrolment.")
        + "</w:endnotes>",
        "word/header1.xml": f"<w:hdr {NS}>{p(r('Running head'))}</w:hdr>",
        "word/footer1.xml": f"<w:ftr {NS}>{p(r('Draft of 14 October'))}</w:ftr>",
        "word/comments.xml": f'<w:comments {NS}><w:comment w:id="0" w:author="a">'
        f"{p(r('This claim is not supported by the data.'))}</w:comment></w:comments>",
    }
    rels = {
        "rIdFn": "footnotes footnotes.xml",
        "rIdEn": "endnotes endnotes.xml",
        "rIdH": "header header1.xml",
        "rIdF": "footer footer1.xml",
        "rIdC": "comments comments.xml",
    }
    return write(path, body, parts, rels)


def test_version_2_reads_the_body_in_document_order_then_the_notes_then_headers_and_footers(
    tmp_path,
):
    assert read(manuscript(tmp_path / "paper.docx")) == [
        "Thirty-five participants responded.",
        "Arm | n",
        "Treated | 1,204",
        "inner | 7",
        "Responders were older by 4.2 years.",
        "A text box beside the figure.",
        "A block inside a content control.",
        "As reported (Smith 2020) and (Jones 2021), see the registry in table 12.",
        "The effect held in every cohort at p<0.05.",
        "A paragraph whose mark is deleted joins the next.",
        "Footnote 1. Participants were paid for their time.",
        "Endnote 1. Age at enrolment.",
        "Running head",
        "Draft of 14 October",
    ]


def test_an_undeclared_docx_is_read_by_version_2(tmp_path):
    r_ = V.check_one("held in every cohort", manuscript(tmp_path / "paper.docx"))
    assert (r_.state, r_.extractor) == ("found", "docx-text@2")


@pytest.mark.parametrize(
    ("passage", "state"),
    [
        ("The effect held in every cohort at", "found"),
        ("The effect vanished at", "not found"),
        ("Footnote 1. Participants were paid for their time.", "found"),
        ("Participants were paid for their time.", "found"),
        ("This claim is not supported by the data.", "not found"),
        ("ZOTERO_ITEM CSL_CITATION", "not found"),
        ("Withdrawn | 3", "not found"),
        ("Running head", "found"),
    ],
)
def test_what_a_quotation_resolves_against(tmp_path, passage, state):
    assert V.check_one(passage, manuscript(tmp_path / "paper.docx")).state == state


def test_a_header_two_sections_share_is_read_once(tmp_path):
    assert read(manuscript(tmp_path / "paper.docx")).count("Running head") == 1


def test_strict_namespaces_are_read_as_transitional_ones(tmp_path):
    strict = "http://purl.oclc.org/ooxml/wordprocessingml/main"
    path = tmp_path / "strict.docx"
    with zipfile.ZipFile(path, "w") as z:
        z.writestr(
            "word/document.xml",
            f'<w:document xmlns:w="{strict}"><w:body>{p(r("Strict text."))}</w:body></w:document>',
        )
    assert read(path) == ["Strict text."]


# --- a source pinned under version 1 ------------------------------------------------------------


def test_a_source_pinned_under_version_1_still_verifies_with_version_1(tmp_path, capsys):
    document = docx.Document()
    document.add_paragraph("Before the table.")
    document.add_table(rows=1, cols=2).rows[0].cells[0].text = "Cohort"
    document.add_paragraph("After the table.")
    source = tmp_path / "sources" / "supp.docx"
    source.parent.mkdir()
    document.save(str(source))
    v1 = E.declared(source, "docx-text@1").read(source)
    # The two versions read these bytes differently, so a v1 pin read by v2 would fail.
    assert v1 != E.Declared(E.EXTRACTORS["docx-text"]).read(source)

    claims = tmp_path / "claims"
    claims.mkdir()
    (claims / "supp.yaml").write_text(
        yaml.safe_dump(
            {
                "source": {
                    "local": "sources/supp.docx",
                    "sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
                    "extractor": "docx-text",
                    "extractor_version": 1,
                    "derived_sha256": hashlib.sha256(v1.encode()).hexdigest(),
                },
                "claims": {"c": {"quotes": [{"exact": "After the table."}]}},
            }
        )
    )
    assert load_claim_file(claims / "supp.yaml").source.reader == "docx-text@1"
    assert cli.main(["verify", "--claims", str(claims), "--strict"]) == 0
    out = capsys.readouterr().out
    assert "docx-text@1" in out and "all found." in out
    assert "no longer extract" not in out


def test_a_version_no_build_ships_names_every_version_this_one_does(tmp_path):
    r_ = V.check_one("held", manuscript(tmp_path / "paper.docx"), None, "docx-text@3")
    assert r_.state == "unchecked"
    assert "ships versions 1 and 2. Each version produces different text" in r_.detail


# --- a hostile file -----------------------------------------------------------------------------


def test_a_part_declaring_entities_is_refused_unread(tmp_path):
    bomb = (
        '<?xml version="1.0"?><!DOCTYPE lolz [<!ENTITY lol "lol">'
        '<!ENTITY lol2 "&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;">]>'
        f'<w:document xmlns:w="{W}"><w:body><w:p><w:r><w:t>&lol2;</w:t></w:r></w:p></w:body></w:document>'
    )
    path = tmp_path / "bomb.docx"
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("word/document.xml", bomb)
    r_ = V.check_one("lol", path)
    assert r_.state == "unchecked"
    assert "declares a document type" in r_.detail


def test_a_package_whose_xml_inflates_past_the_cap_is_refused_unread(tmp_path, monkeypatch):
    path = write(tmp_path / "big.docx", p(r("x" * 5000)))
    monkeypatch.setattr(E, "DOCX_MAX_XML_BYTES", 1000)
    r_ = V.check_one("xxxx", path)
    assert r_.state == "unchecked"
    assert "inflates past 1,000 bytes" in r_.detail


def test_a_compressed_part_is_capped_on_what_it_inflates_to_not_what_it_claims(
    tmp_path, monkeypatch
):
    path = tmp_path / "bomb.docx"
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as z:
        z.writestr(
            "word/document.xml",
            f'<w:document xmlns:w="{W}"><w:body>' + " " * 2_000_000 + "</w:body></w:document>",
        )
    assert path.stat().st_size < 10_000
    monkeypatch.setattr(E, "DOCX_MAX_XML_BYTES", 1_000_000)
    assert "inflates past" in V.check_one("anything at all", path).detail


@pytest.mark.parametrize(
    ("body", "reason"),
    [
        (b"not a zip archive at all", "not a zip archive"),
        (None, "holds no word/document.xml"),
    ],
)
def test_a_file_that_is_not_a_docx_is_unchecked_with_the_reason(tmp_path, body, reason):
    path = tmp_path / "fake.docx"
    if body is None:
        with zipfile.ZipFile(path, "w") as z:
            z.writestr("other.xml", "<a/>")
    else:
        path.write_bytes(body)
    r_ = V.check_one("anything at all", path)
    assert r_.state == "unchecked"
    assert reason in r_.detail


# --- what Word writes that a first reading got wrong -----------------------------------------


def field(instruction: str, result: str) -> str:
    """A complex field, as Word writes one: begin, instruction, separate, result, end."""
    return (
        '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
        f'<w:r><w:instrText xml:space="preserve"> {instruction} </w:instrText></w:r>'
        '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
        + r(result)
        + '<w:r><w:fldChar w:fldCharType="end"/></w:r>'
    )


def test_a_page_number_field_is_not_read_and_its_neighbours_are(tmp_path):
    footer = p(
        r("Page "),
        field("PAGE \\* MERGEFORMAT", "3"),
        r(" of "),
        f'<w:fldSimple w:instr=" NUMPAGES ">{r("17")}</w:fldSimple>',
        r(", section page "),
        field("SECTIONPAGES", "4"),
    )
    body = p(r("See the methods on page "), field("PAGEREF _Ref1 \\h", "9"), r(".")) + (
        '<w:sectPr><w:footerReference w:type="default" r:id="rIdF"/></w:sectPr>'
    )
    path = write(
        tmp_path / "paper.docx",
        body,
        {"word/footer1.xml": f"<w:ftr {NS}>{footer}</w:ftr>"},
        {"rIdF": "footer footer1.xml"},
    )
    # The page a sentence refers the reader to is the author's; the footer's are the layout's.
    assert read(path) == ["See the methods on page 9.", "Page of , section page"]


def test_a_field_that_never_ends_hides_nothing_after_its_paragraph(tmp_path):
    unended = '<w:r><w:fldChar w:fldCharType="begin"/></w:r>' + (
        "<w:r><w:instrText>REF x</w:instrText></w:r>"
    )
    end_deleted = (
        '<w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText>REF x</w:instrText></w:r>'
        '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
        + r("res")
        + '<w:del w:id="1" w:author="a"><w:r><w:fldChar w:fldCharType="end"/></w:r></w:del>'
    )
    body = p(r("Before "), unended) + p(r("following paragraph")) + p(end_deleted) + p(r("last"))
    assert read(write(tmp_path / "paper.docx", body)) == [
        "Before",
        "following paragraph",
        "res",
        "last",
    ]


def test_text_a_change_deletes_inside_an_equation_is_not_read(tmp_path):
    math = (
        "<m:oMath><m:r><m:t>a</m:t></m:r>"
        '<w:del w:id="1" w:author="a"><m:r><m:t>DELETED</m:t></m:r></w:del>'
        '<m:r><w:rPr><w:del w:id="2" w:author="a"/></w:rPr><m:t>GONE</m:t></m:r>'
        '<w:ins w:id="3" w:author="a"><m:r><m:t>b</m:t></m:r></w:ins></m:oMath>'
    )
    assert read(write(tmp_path / "paper.docx", p(r("Eq: "), math))) == ["Eq: ab"]


def test_a_paragraph_whose_mark_is_deleted_before_a_table_stays_out_of_the_table(tmp_path):
    body = (
        p(r("first half"), mark='<w:pPr><w:rPr><w:del w:id="1" w:author="a"/></w:rPr></w:pPr>')
        + table(row(cell(p(r("cell")))))
        + p(r("second half"))
    )
    assert read(write(tmp_path / "paper.docx", body)) == ["first half", "cell", "second half"]


@pytest.mark.parametrize(
    ("font", "code", "shown"),
    [
        ("Symbol", "F0A3", "≤"),
        ("Symbol", "F0B3", "≥"),
        ("Symbol", "F0B1", "±"),
        ("Symbol", "F0B4", "×"),
        ("Symbol", "F0B8", "÷"),
        ("Symbol", "F0B9", "≠"),
        ("Symbol", "F0BB", "≈"),
        ("Symbol", "F06D", "μ"),
        ("Symbol", "F061", "α"),
        ("Symbol", "F044", "Δ"),
        ("Wingdings", "F0FC", "✓"),
        ("Times New Roman", "00B5", "µ"),
    ],
)
def test_a_symbol_is_read_as_the_character_it_draws(tmp_path, font, code, shown):
    sym = f'<w:r><w:sym w:font="{font}" w:char="{code}"/></w:r>'
    assert read(write(tmp_path / "paper.docx", p(r("p "), sym, r(" 0.05")))) == [f"p {shown} 0.05"]


def test_a_ruby_is_read_as_its_base_text_and_not_its_reading(tmp_path):
    ruby = (
        "<w:r><w:ruby><w:rubyPr/><w:rt>"
        + r("kan")
        + "</w:rt><w:rubyBase>"
        + r("BASE")
        + "</w:rubyBase></w:ruby></w:r>"
    )
    assert read(write(tmp_path / "paper.docx", p(r("See "), ruby, r(" here")))) == ["See BASE here"]


@pytest.mark.parametrize("encoding", ["utf-16-le", "utf-16-be", "utf-8-sig"])
def test_a_part_is_decoded_by_its_byte_order_mark(tmp_path, encoding):
    xml = f'<?xml version="1.0" encoding="{"UTF-8" if encoding == "utf-8-sig" else "UTF-16"}"?>'
    xml += f'<w:document xmlns:w="{W}"><w:body>{p(r("Read in either encoding."))}</w:body></w:document>'
    mark = {"utf-16-le": b"\xff\xfe", "utf-16-be": b"\xfe\xff", "utf-8-sig": b""}[encoding]
    path = tmp_path / "encoded.docx"
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("word/document.xml", mark + xml.encode(encoding))
    assert read(path) == ["Read in either encoding."]


def test_a_utf16_part_declaring_entities_is_still_refused(tmp_path):
    xml = f'<?xml version="1.0" encoding="UTF-16"?><!DOCTYPE d [<!ENTITY e "x">]><w:document xmlns:w="{W}"/>'
    path = tmp_path / "bomb.docx"
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("word/document.xml", xml.encode("utf-16"))
    assert "declares a document type" in V.check_one("anything at all", path).detail


# --- a source pinned before versions were recorded -----------------------------------------------


def toc_document(path):
    """A table of contents, which version 1 reads without its field result's page numbers in
    place and version 2 reads as Word shows it: the two readings differ."""
    entry = p(field("TOC \\o", "Methods"), r(" and results"))
    return write(path, p(r("Contents")) + entry + table(row(cell(p(r("n"))))) + p(r("Methods.")))


def pinned_claims(tmp_path, quotes, **source):
    claims = tmp_path / "claims"
    claims.mkdir(exist_ok=True)
    path = claims / "supp.yaml"
    path.write_text(
        yaml.safe_dump(
            {
                "source": {"local": "sources/supp.docx"} | source,
                "claims": {f"c{i}": {"quotes": [{"exact": q}]} for i, q in enumerate(quotes)},
            }
        )
    )
    return path


@pytest.mark.parametrize("declared", [True, False])
def test_a_docx_source_recording_no_version_is_read_with_version_1(tmp_path, capsys, declared):
    (tmp_path / "sources").mkdir()
    source = toc_document(tmp_path / "sources" / "supp.docx")
    v1 = E.declared(source, "docx-text@1").read(source)
    assert v1 != E.Declared(E.EXTRACTORS["docx-text"]).read(source)
    named = {"extractor": "docx-text"} if declared else {}
    path = pinned_claims(
        tmp_path,
        ["Contents"],
        sha256=hashlib.sha256(source.read_bytes()).hexdigest(),
        derived_sha256=hashlib.sha256(v1.encode()).hexdigest(),
        **named,
    )
    assert load_claim_file(path).source.reader == "docx-text@1"
    assert cli.main(["verify", "--claims", str(path.parent), "--strict"]) == 0
    out = capsys.readouterr().out
    assert "docx-text@1" in out and "no longer extract" not in out
    assert (
        "supp: read with docx-text@1, inferred: no version recorded; set extractor_version to "
        "choose" in out
    )


def test_a_new_pin_in_a_docx_records_version_2_and_its_text(tmp_path):
    (tmp_path / "sources").mkdir()
    source = toc_document(tmp_path / "sources" / "supp.docx")
    path = pinned_claims(tmp_path, [], sha256=hashlib.sha256(source.read_bytes()).hexdigest())
    assert pin.main([str(path), "--id", "toc", "--quote", "Methods and results"]) == 0
    written = load_claim_file(path).source
    v2 = E.Declared(E.EXTRACTORS["docx-text"]).read(source)
    assert (written.extractor, written.extractor_version) == ("docx-text", 2)
    assert written.derived_sha256 == hashlib.sha256(v2.encode()).hexdigest()
    assert written.reader == "docx-text@2"


def test_a_page_reference_inside_a_table_of_contents_is_not_read(tmp_path):
    entry = p(r("Methods "), field("PAGEREF _Toc1 \\h", "3"))
    styled = p(
        r("Results "),
        field("PAGEREF _Toc2 \\h", "5"),
        mark='<w:pPr><w:pStyle w:val="TOC1"/></w:pPr>',
    )
    contents = (
        p(
            '<w:r><w:fldChar w:fldCharType="begin"/></w:r>',
            '<w:r><w:instrText> TOC \\o "1-3" </w:instrText></w:r>',
            '<w:r><w:fldChar w:fldCharType="separate"/></w:r>',
            r("Introduction "),
            field("PAGEREF _Toc0 \\h", "1"),
        )
        + entry
        + p('<w:r><w:fldChar w:fldCharType="end"/></w:r>')
    )
    body = contents + styled + p(r("As shown on page "), field("PAGEREF _Toc1 \\h", "3"), r("."))
    assert read(write(tmp_path / "paper.docx", body)) == [
        "Introduction",
        "Methods",
        "Results",
        "As shown on page 3.",
    ]


def test_a_formula_over_the_page_number_is_not_read(tmp_path):
    nested = (
        '<w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText>= </w:instrText></w:r>'
        + field("PAGE", "4")
        + '<w:r><w:instrText> - 1</w:instrText></w:r><w:r><w:fldChar w:fldCharType="separate"/></w:r>'
        + r("3")
        + '<w:r><w:fldChar w:fldCharType="end"/></w:r>'
    )
    simple = '<w:fldSimple w:instr=" = NUMPAGES - 2 ">' + r("15") + "</w:fldSimple>"
    sums = '<w:fldSimple w:instr=" = 2 + 2 ">' + r("4") + "</w:fldSimple>"
    body = p(r("Page "), nested, r(" of "), simple, r(", and "), sums, r("."))
    assert read(write(tmp_path / "paper.docx", body)) == ["Page of , and 4."]


@pytest.mark.parametrize("encoding", ["utf-16-le", "utf-16-be"])
def test_a_utf16_part_with_no_byte_order_mark_declaring_entities_is_refused(tmp_path, encoding):
    entities = '<!ENTITY a "AAAAAAAAAA">' + "".join(
        f'<!ENTITY {chr(98 + i)} "' + f"&{chr(97 + i)};" * 10 + '">' for i in range(4)
    )
    xml = (
        f'<?xml version="1.0" encoding="UTF-16"?><!DOCTYPE w:document [{entities}]>'
        f'<w:document xmlns:w="{W}"><w:body>{p(r("&e;"))}</w:body></w:document>'
    )
    path = tmp_path / "bomb.docx"
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("word/document.xml", xml.encode(encoding))
    r_ = V.check_one("AAAAAAAAAA", path)
    assert r_.state == "unchecked"
    assert "declares a document type" in r_.detail


@pytest.mark.parametrize(
    ("declared", "encoding"),
    [
        ("UTF-16", "utf-16-le"),
        ("UTF-16", "utf-16-be"),
        ("ISO-8859-1", "latin-1"),
        ("windows-1252", "cp1252"),
        ("us-ascii", "ascii"),
    ],
)
def test_a_part_is_decoded_as_the_parser_detects_its_encoding(tmp_path, declared, encoding):
    text = "Na\u00efve readers \u00e9crit." if encoding != "ascii" else "Plain readers."
    body = p(r(text))
    xml = (
        f'<?xml version="1.0" encoding="{declared}"?>'
        f'<w:document xmlns:w="{W}"><w:body>{body}</w:body></w:document>'
    )
    path = tmp_path / "encoded.docx"
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("word/document.xml", xml.encode(encoding))
    assert read(path) == [text]


def test_a_part_holding_a_nul_is_refused(tmp_path):
    path = tmp_path / "nul.docx"
    with zipfile.ZipFile(path, "w") as z:
        z.writestr(
            "word/document.xml", f'<w:document xmlns:w="{W}"><w:body>\x00</w:body></w:document>'
        )
    assert "holds a NUL" in V.check_one("anything at all", path).detail


def test_a_source_written_by_hand_with_no_version_and_no_digest_is_read_with_version_2(
    tmp_path, capsys
):
    (tmp_path / "sources").mkdir()
    source = manuscript(tmp_path / "sources" / "supp.docx")
    footnote = "Participants were paid for their time."
    path = pinned_claims(
        tmp_path,
        [footnote],
        sha256=hashlib.sha256(source.read_bytes()).hexdigest(),
        extractor="docx-text",
    )
    cf = load_claim_file(path)
    assert (cf.source.reader, cf.version_inferred) == ("docx-text", "")
    assert cli.main(["verify", "--claims", str(path.parent), "--strict"]) == 0
    out = capsys.readouterr().out
    assert "docx-text@2" in out and "inferred" not in out
    assert (
        pin.main([str(path), "--id", "second", "--quote", "held in every cohort", "--check"]) == 0
    )


@pytest.mark.parametrize("declared", ["UTF-7", "unicode_escape", "cp037", "no-such-codec"])
def test_a_part_declaring_an_encoding_no_word_part_is_in_is_refused_by_name(tmp_path, declared):
    # `unicode_escape` turns an escaped `<!DOCTYPE` in plain ASCII into a declaration; the
    # encoding is refused by name before anything is decoded with it.
    xml = f'<?xml version="1.0" encoding="{declared}"?><w:document xmlns:w="{W}"/>'
    path = tmp_path / "odd.docx"
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("word/document.xml", xml.encode("ascii"))
    r_ = V.check_one("anything at all", path)
    assert r_.state == "unchecked"
    assert f"declares {declared}" in r_.detail
