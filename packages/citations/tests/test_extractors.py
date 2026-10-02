"""A source that is a workbook, a `.docx` or an article's XML is pinned as that file.

Before these extractors a claims file could pin such a source only as a text somebody's script
had made from it, and a reader who downloaded the publisher's file held bytes no pin named. The
properties under test: a quotation resolves in the original through a named, versioned
extractor; a reader that is not installed, a version this build does not ship, and an option an
extractor does not take are each `unchecked` and never `found`; and `fetch` verifies the
original first and the extracted text second, reporting a failure at each stage as its own
outcome.
"""

from __future__ import annotations

import dataclasses
import datetime
import hashlib

import docx
import openpyxl
import pytest
import xlwt
import yaml
from citations import cli, pin
from citations import extractors as E
from citations import fetch as F
from citations import verify as V
from citations.exceptions import ClaimFileError
from citations.models import ClaimSource, load_claim_file

HEADER = ["Cohort", "Cases", "Controls", "Mean age"]
UKB = ["UK Biobank", 34541, 261984, 56.5]
DECODE = ["deCODE", 37379, 318845, 54.0]
UKB_ROW = "UK Biobank | 34541 | 261984 | 56.5"

ARTICLE = """<?xml version="1.0" encoding="UTF-8"?>
<article xmlns:xlink="http://www.w3.org/1999/xlink">
  <front><article-meta><title-group>
    <article-title>Discovery of 38 loci for <italic>coronary</italic> disease</article-title>
  </title-group></article-meta></front>
  <body>
    <sec>
      <title>Methods</title>
      <p>We analyzed <italic>APOE</italic> in 86,847 cases<sup><xref rid="r1">1</xref></sup>
      &amp; 417,789 controls at &#x3b1; = 0.05 (<xref rid="t1">Table 1</xref>).</p>
      <table-wrap><table>
        <thead><tr><th>Cohort</th><th>Cases</th></tr></thead>
        <tbody><tr><td>UK <bold>Biobank</bold></td><td>34,541</td></tr></tbody>
      </table></table-wrap>
    </sec>
  </body>
</article>
"""
SENTENCE = "We analyzed APOE in 86,847 cases1 & 417,789 controls at α = 0.05 (Table 1)."


@pytest.fixture(autouse=True)
def _no_cache():
    V.clear_caches()


def workbook(path, sheets):
    book = openpyxl.Workbook()
    book.remove(book.active)
    for name, rows in sheets.items():
        ws = book.create_sheet(name)
        for row in rows:
            ws.append(row)
    book.save(path)
    return path


def article(tmp_path, text=ARTICLE):
    path = tmp_path / "article.xml"
    path.write_text(text, encoding="utf-8")
    return path


def document(path):
    doc = docx.Document()
    doc.add_paragraph("Supplementary Note. Cohorts were genotyped on the OncoArray.")
    table = doc.add_table(rows=3, cols=3)
    table.cell(0, 0).merge(table.cell(0, 2)).text = "Discovery studies"
    for r, row in enumerate([["BCAC", "61,282", "45,494"], ["CIMBA", "", "9,414"]], start=1):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    doc.add_paragraph("Counts are after quality control.")
    doc.save(str(path))
    return path


# --- a row of a workbook ------------------------------------------------------------------------


def test_a_row_quoted_from_a_workbook_resolves_in_the_workbook(tmp_path):
    source = workbook(tmp_path / "supp.xlsx", {"ST1": [HEADER, UKB, DECODE]})
    r = V.check_one(UKB_ROW, source)
    assert r.state == "found"
    assert r.extractor == "sheet-rows@1", "the result names the extractor and its version"
    assert (
        r.extraction_digest
        == hashlib.sha256(E.Declared(E.EXTRACTORS["sheet-rows"]).read(source).encode()).hexdigest()
    )


def test_cells_taken_from_two_rows_are_not_a_row_of_the_workbook(tmp_path):
    # The source was read and the passage is not in it: `not found`, which only a reader that
    # ran can say. `unchecked` here would mean the workbook was never opened.
    source = workbook(tmp_path / "supp.xlsx", {"ST1": [HEADER, UKB, DECODE]})
    r = V.check_one("UK Biobank | 37379 | 318845", source)
    assert r.state == "not found"
    assert r.extractor == "sheet-rows@1"


def test_one_row_is_one_line_and_cells_are_in_column_order(tmp_path):
    source = workbook(tmp_path / "supp.xlsx", {"ST1": [HEADER, UKB, DECODE]})
    text = E.Declared(E.EXTRACTORS["sheet-rows"]).read(source)
    assert text.splitlines() == [
        "Cohort | Cases | Controls | Mean age",
        UKB_ROW,
        "deCODE | 37379 | 318845 | 54",
    ]


def test_a_named_sheet_is_the_only_sheet_read(tmp_path):
    source = workbook(
        tmp_path / "supp.xlsx", {"Contents": [["Table", "Title"]], "ST1": [HEADER, UKB]}
    )
    assert V.check_one(UKB_ROW, source, None, "sheet-rows --sheet ST1").state == "found"
    elsewhere = V.check_one(UKB_ROW, source, None, "sheet-rows --sheet Contents")
    assert elsewhere.state == "not found"
    assert elsewhere.extractor == "sheet-rows@1 --sheet Contents"


def test_a_row_on_two_sheets_is_ambiguous_until_the_sheet_is_named(tmp_path):
    # A workbook repeats a cohort row on a summary sheet. Read whole, the quotation points at
    # two rows and identifies neither; the sheet is what says which was meant.
    source = workbook(tmp_path / "supp.xlsx", {"Summary": [UKB], "ST1": [HEADER, UKB]})
    assert V.check_one(UKB_ROW, source).state == "ambiguous"
    assert V.check_one(UKB_ROW, source, None, "sheet-rows --sheet ST1").state == "found"


def test_a_sheet_the_workbook_does_not_hold_is_unchecked_and_names_the_ones_it_does(tmp_path):
    source = workbook(tmp_path / "supp.xlsx", {"ST1": [HEADER, UKB]})
    r = V.check_one(UKB_ROW, source, None, "sheet-rows --sheet 'Table S1'")
    assert r.state == "unchecked"
    assert "'Table S1'" in r.detail and "'ST1'" in r.detail


def test_empty_cells_are_dropped_unless_the_source_keeps_them(tmp_path):
    source = workbook(tmp_path / "supp.xlsx", {"ST1": [["FinnGen", None, 1204, None, None]]})
    dropped = "FinnGen | 1204"
    kept = "FinnGen | | 1204"
    assert V.check_one(dropped, source).state == "found"
    assert V.check_one(kept, source).state == "not found"
    assert V.check_one(kept, source, None, "sheet-rows --keep-empty").state == "found"
    assert V.check_one(dropped, source, None, "sheet-rows --keep-empty").state == "not found"


def test_a_number_reads_as_the_sheet_shows_it_and_not_as_the_float_stored(tmp_path):
    # 0.1 + 0.2 is stored as 0.30000000000000004 and shown as 0.3; 46 is stored as 46.0.
    source = workbook(
        tmp_path / "supp.xlsx",
        {"ST1": [["age", 46.0, 0.1 + 0.2, True, datetime.datetime(2021, 3, 9)]]},
    )
    assert E.Declared(E.EXTRACTORS["sheet-rows"]).read(source) == (
        "age | 46 | 0.3 | TRUE | 2021-03-09"
    )


def test_a_line_break_inside_a_cell_does_not_start_a_new_row(tmp_path):
    source = workbook(tmp_path / "supp.xlsx", {"ST1": [["Rotterdam\nStudy", 4521]]})
    assert E.Declared(E.EXTRACTORS["sheet-rows"]).read(source) == "Rotterdam Study | 4521"


def test_a_csv_is_read_as_rows_without_its_byte_order_mark(tmp_path):
    source = tmp_path / "cohorts.csv"
    source.write_bytes(
        b'\xef\xbb\xbfCohort,Cases,Controls\r\n"Biobank Japan, BBJ",29319,183134\r\n'
    )
    r = V.check_one("Biobank Japan, BBJ | 29319 | 183134", source)
    assert r.state == "found"
    assert E.Declared(E.EXTRACTORS["sheet-rows"]).read(source).startswith("Cohort | Cases")
    named = V.check_one("Cohort | Cases", source, None, "sheet-rows --sheet ST1")
    assert named.state == "unchecked" and "no sheets" in named.detail


def test_an_xls_workbook_is_read_with_its_dates_and_whole_numbers(tmp_path):
    # The old format stores every number as a float and a date as a float tagged as one.
    source = tmp_path / "supp.xls"
    book = xlwt.Workbook()
    ws = book.add_sheet("ST1")
    dated = xlwt.easyxf(num_format_str="YYYY-MM-DD")
    for c, value in enumerate(UKB):
        ws.write(0, c, value)
    ws.write(1, 0, "frozen")
    ws.write(1, 1, datetime.datetime(2021, 3, 9), dated)
    book.save(str(source))

    assert V.check_one(UKB_ROW, source).state == "found"
    assert V.check_one("frozen | 2021-03-09", source).state == "found"
    assert V.check_one(UKB_ROW, source, None, "sheet-rows --sheet ST9").state == "unchecked"


@pytest.mark.parametrize(("module", "suffix"), [("openpyxl", ".xlsx"), ("xlrd", ".xls")])
def test_a_workbook_with_no_reader_installed_is_unchecked_and_names_the_install(
    tmp_path, monkeypatch, module, suffix
):
    # The outcome for a source nothing here could read. `found` is unreachable, and so is
    # `not found`: an uninstalled parser says nothing about what the workbook holds.
    source = tmp_path / f"supp{suffix}"
    source.write_bytes(b"not opened: the reader is absent")
    monkeypatch.setattr(E, module, None)
    r = V.check_one(UKB_ROW, source)
    assert r.state == "unchecked"
    assert f"{module} is not installed" in r.detail and "citations[sheets]" in r.detail
    assert r.extractor == "", "nothing read it, so nothing is named"


def test_a_file_that_is_not_a_workbook_is_unchecked_not_empty(tmp_path):
    source = tmp_path / "supp.xlsx"
    source.write_bytes(b"<html>a landing page saved under the supplement's name</html>")
    r = V.check_one(UKB_ROW, source)
    assert r.state == "unchecked"
    assert "sheet-rows could not read it" in r.detail


def test_a_workbook_is_not_handed_to_the_pdf_readers_when_a_row_is_absent(tmp_path, monkeypatch):
    # `not found` on a PDF asks the other PDF readers before accusing the manuscript. On a
    # workbook that would run four programs that cannot open it and name them as having looked.
    def never(*a, **kw):
        raise AssertionError("a PDF reader was pointed at a workbook")

    monkeypatch.setattr(V, "_poppler", never)
    monkeypatch.setattr(V.subprocess, "run", never)
    source = workbook(tmp_path / "supp.xlsx", {"ST1": [HEADER, UKB]})
    r = V.check_one("a row that is nowhere in this workbook | 1 | 2", source, page=3)
    assert r.state == "not found"
    assert "pdftotext" not in r.detail
    assert "page" not in V.check_one(UKB_ROW, source, page=3).warnings
    # Nor triangulated among them: every one would fail to open it, and a row that is there
    # would come back `unchecked`.
    assert V.check_one(UKB_ROW, source, triangulate=True).state == "found"


# --- a .docx --------------------------------------------------------------------------------------


def test_a_paragraph_and_a_table_row_of_a_docx_both_resolve(tmp_path):
    source = document(tmp_path / "note.docx")
    paragraph = V.check_one("Cohorts were genotyped on the OncoArray.", source)
    row = V.check_one("BCAC | 61,282 | 45,494", source)
    assert paragraph.state == row.state == "found"
    assert row.extractor == "docx-text@1"
    assert V.check_one("BCAC | 9,414", source).state == "not found"


def test_a_docx_reads_as_its_paragraphs_then_its_table_rows(tmp_path):
    # The merged heading spans three columns and is returned once per column; printed once.
    assert E.Declared(E.EXTRACTORS["docx-text"]).read(
        document(tmp_path / "note.docx")
    ).splitlines() == [
        "Supplementary Note. Cohorts were genotyped on the OncoArray.",
        "Counts are after quality control.",
        "Discovery studies",
        "BCAC | 61,282 | 45,494",
        "CIMBA | 9,414",
    ]


def test_a_docx_with_no_reader_installed_is_unchecked_and_names_the_install(tmp_path, monkeypatch):
    source = document(tmp_path / "note.docx")
    monkeypatch.setattr(E, "docx", None)
    r = V.check_one("Cohorts were genotyped on the OncoArray.", source)
    assert r.state == "unchecked"
    assert "python-docx is not installed" in r.detail and "citations[docx]" in r.detail


# --- an article's XML -----------------------------------------------------------------------------


def test_a_sentence_crossing_inline_tags_resolves_once_they_are_stripped(tmp_path):
    # The before and the after on one file. Read as text the XML keeps `<italic>`, `<sup>`,
    # `<xref>` and `&amp;`, and the sentence as printed is not in it.
    source = article(tmp_path)
    assert V.check_one(SENTENCE, source).state == "not found"
    r = V.check_one(SENTENCE, source, None, "jats-text")
    assert r.state == "found"
    assert r.extractor == "jats-text@1"


def test_each_block_is_a_line_and_a_table_row_is_its_cells(tmp_path):
    assert E.Declared(E.EXTRACTORS["jats-text"]).read(article(tmp_path)).splitlines() == [
        "Discovery of 38 loci for coronary disease",
        "Methods",
        SENTENCE,
        "Cohort | Cases",
        "UK Biobank | 34,541",
    ]


def test_an_element_the_extractor_does_not_know_keeps_its_text(tmp_path):
    # Blocks are whatever is not listed as inline, so an unfamiliar element costs a line break
    # and never its words. A list of blocks would have dropped this sentence.
    source = article(
        tmp_path,
        "<article><body><p>Before the box.</p><boxed-text><unheard-of>The effect was "
        "replicated in an independent sample of 5,346 participants.</unheard-of></boxed-text>"
        "</body></article>",
    )
    quote = "The effect was replicated in an independent sample of 5,346 participants."
    assert V.check_one(quote, source, None, "jats-text").state == "found"


def test_xml_that_is_not_well_formed_is_unchecked(tmp_path):
    source = article(tmp_path, "<article><body><p>An unclosed paragraph.</body></article>")
    r = V.check_one("An unclosed paragraph.", source, None, "jats-text")
    assert r.state == "unchecked"
    assert "not well-formed XML" in r.detail


def test_an_xml_source_declaring_nothing_is_still_read_with_its_markup(tmp_path):
    # Quotations already pinned against the XML as served must go on resolving: the suffix
    # does not choose `jats-text`, a source names it.
    source = article(tmp_path)
    r = V.check_one("We analyzed <italic>APOE</italic> in 86,847 cases", source)
    assert r.state == "found"
    assert r.extractor == V.PLAIN_TEXT
    assert ".xml" not in E.BY_SUFFIX


# --- named and versioned -------------------------------------------------------------------------


def test_a_built_in_extractor_needs_no_consent_and_runs_no_program(tmp_path, monkeypatch):
    def never(*a, **kw):
        raise AssertionError("a built-in extractor must not reach subprocess")

    monkeypatch.setattr(V.subprocess, "run", never)
    r = V.check_one(SENTENCE, article(tmp_path), None, "jats-text", frozenset())
    assert r.state == "found"


def test_a_version_this_build_does_not_ship_is_unchecked_and_names_both(tmp_path):
    # Reading with the version that is here would report a verdict against text the claims
    # file did not pin its quotations in.
    source = article(tmp_path)
    assert V.check_one(SENTENCE, source, None, "jats-text@1").state == "found"
    r = V.check_one(SENTENCE, source, None, "jats-text@2")
    assert r.state == "unchecked"
    assert "version 2" in r.detail and "ships version 1" in r.detail
    assert r.extractor == ""


def test_an_option_an_extractor_does_not_take_is_unchecked(tmp_path):
    r = V.check_one(SENTENCE, article(tmp_path), None, "jats-text --sheet ST1")
    assert r.state == "unchecked"
    assert "jats-text does not take '--sheet'" in r.detail


def test_every_suffix_names_an_extractor_that_exists():
    assert set(E.BY_SUFFIX.values()) <= set(E.EXTRACTORS)
    assert all(name == e.name and e.version >= 1 for name, e in E.EXTRACTORS.items())


def test_the_source_block_reaches_the_extractor_with_its_version_and_options():
    source = ClaimSource(
        extractor="sheet-rows", extractor_version=1, sheet="S1-Cohort Table", empty_cells="keep"
    )
    got = E.declared(None, source.reader)
    assert got == E.Declared(E.EXTRACTORS["sheet-rows"], "S1-Cohort Table", True)
    assert ClaimSource(extract_cmd="detex").reader == "detex"
    assert E.declared(None, "detex") is None, "a program is not a built-in extractor"


def test_a_source_naming_an_extractor_and_a_command_is_refused(tmp_path):
    path = tmp_path / "both.yaml"
    path.write_text("source:\n  extractor: jats-text\n  extract_cmd: detex\nclaims: {}\n")
    with pytest.raises(ClaimFileError, match="both `extractor` and `extract_cmd`"):
        load_claim_file(path)


# --- through the claims file: the original, the extractor, the derived text ---------------------


def derived_of(source, reader) -> str:
    return hashlib.sha256(V.extract_uncached(source, None, reader).encode()).hexdigest()


def claims_for(tmp_path, body: bytes, quotes=(), local="sources/article.xml", **fields):
    """A claims directory pinning `body`, which is not on disk until a test puts it there."""
    claims = tmp_path / "claims"
    claims.mkdir(exist_ok=True)
    block = {"local": local, "sha256": hashlib.sha256(body).hexdigest()} | fields
    path = claims / "article.yaml"
    path.write_text(
        yaml.safe_dump(
            {
                "source": block,
                "claims": {f"c{i}": {"quotes": [{"exact": q}]} for i, q in enumerate(quotes)},
            },
            allow_unicode=True,
        )
    )
    return load_claim_file(path)


def on_disk(cf, body: bytes):
    cf.artifact().parent.mkdir(parents=True, exist_ok=True)
    cf.artifact().write_bytes(body)
    return cf.artifact()


def changed(monkeypatch, name="jats-text"):
    """The same extractor under the same version, now producing other text from the same bytes."""
    was = E.EXTRACTORS[name]
    monkeypatch.setitem(
        E.EXTRACTORS,
        name,
        dataclasses.replace(was, read=lambda *a: was.read(*a).replace("Methods", "METHODS")),
    )


def test_verify_resolves_a_quotation_in_the_original_through_the_named_extractor(tmp_path, capsys):
    body = ARTICLE.encode()
    source = on_disk(claims_for(tmp_path, body), body)
    cf = claims_for(
        tmp_path,
        body,
        [SENTENCE, "UK Biobank | 34,541"],
        extractor="jats-text",
        extractor_version=1,
        derived_sha256=derived_of(source, "jats-text"),
    )
    code = cli.main(["verify", "--claims", str(cf.path.parent), "--strict"])
    out = capsys.readouterr().out
    assert code == 0
    assert "jats-text@1" in out and "all found." in out


def test_verify_says_so_when_the_extractor_no_longer_produces_the_pinned_text(
    tmp_path, capsys, monkeypatch
):
    # The pin holds: the bytes did not move. The reading of them did, and a quotation that
    # still resolves is resolving in text the record does not describe.
    body = ARTICLE.encode()
    source = on_disk(claims_for(tmp_path, body), body)
    recorded = derived_of(source, "jats-text")
    cf = claims_for(tmp_path, body, [SENTENCE], extractor="jats-text", derived_sha256=recorded)
    changed(monkeypatch)
    V.clear_caches()

    code = cli.main(["verify", "--claims", str(cf.path.parent)])
    out = capsys.readouterr().out
    assert code == 1
    assert "1 source no longer extracts to the text that was pinned" in out
    assert recorded[:12] in out
    assert "changed since being pinned" not in out, "the pin is intact and must not be blamed"


def test_verify_does_not_report_a_changed_reading_for_a_source_nothing_could_read(tmp_path, capsys):
    body = ARTICLE.encode()
    cf = claims_for(
        tmp_path,
        body,
        [SENTENCE],
        extractor="jats-text",
        extractor_version=7,
        derived_sha256="0" * 64,
    )
    on_disk(cf, body)
    cli.main(["verify", "--claims", str(cf.path.parent)])
    out = capsys.readouterr().out
    assert "no longer extract" not in out
    assert "unchecked" in out and "ships version 1" in out


# --- fetch: the original first, the derived text second ------------------------------------------

URL = "https://www.ebi.ac.uk/europepmc/webservices/rest/PMC1/fullTextXML"
BODY = ARTICLE.encode()
LANDING = b"<html>a landing page</html>"


@pytest.fixture
def served(monkeypatch):
    bodies: dict[str, bytes | None] = {}
    monkeypatch.setattr(F, "download", bodies.get)
    monkeypatch.setattr(F.resolve, "get", lambda url, as_json: None)
    return bodies


@pytest.fixture
def pinned(tmp_path):
    """A claims file pinning the article's XML, `jats-text`, and the text that produces."""
    source = on_disk(claims_for(tmp_path, BODY), BODY)
    derived = derived_of(source, "jats-text")
    source.unlink()
    return claims_for(tmp_path, BODY, url=URL, extractor="jats-text", derived_sha256=derived)


def test_fetch_installs_the_original_and_confirms_the_text_extracted_from_it(pinned, served):
    served[URL] = BODY
    outcome = F.fetch_one(pinned)
    assert outcome.state == "fetched"
    assert pinned.artifact().read_bytes() == BODY


def test_fetch_stops_at_the_original_when_its_bytes_differ(pinned, served, monkeypatch):
    # Stage one. Nothing is written, and the extractor is never run on bytes no pin names.
    def never(*a):
        raise AssertionError("the extractor ran on bytes that failed the pin")

    monkeypatch.setitem(
        E.EXTRACTORS, "jats-text", dataclasses.replace(E.EXTRACTORS["jats-text"], read=never)
    )
    served[URL] = LANDING
    outcome = F.fetch_one(pinned)
    assert outcome.state == "differs"
    assert hashlib.sha256(LANDING).hexdigest()[:12] in outcome.detail
    assert not pinned.artifact().exists()


def test_fetch_reports_a_changed_extraction_apart_from_a_changed_file(pinned, served, monkeypatch):
    # Stage two. The bytes are the pinned ones, so they are installed; what failed is the
    # reading, and it is not reported as `differs`, which would send a reader after another
    # copy of a file they already hold.
    served[URL] = BODY
    changed(monkeypatch)
    outcome = F.fetch_one(pinned)
    assert outcome.state == "text differs"
    assert pinned.source.derived_sha256[:12] in outcome.detail
    assert pinned.artifact().read_bytes() == BODY


def test_fetch_does_not_call_a_text_it_could_not_extract_a_match(tmp_path, served, monkeypatch):
    # The original verifies and the reader for it is not installed. `fetched` would claim the
    # second stage passed; it was never run.
    book = workbook(tmp_path / "made.xlsx", {"ST1": [HEADER, UKB]})
    body = book.read_bytes()
    cf = claims_for(
        tmp_path,
        body,
        local="sources/supp.xlsx",
        url=URL,
        extractor="sheet-rows",
        derived_sha256=derived_of(book, "sheet-rows"),
    )
    served[URL] = body
    monkeypatch.setattr(E, "openpyxl", None)
    outcome = F.fetch_one(cf)
    assert outcome.state == "text unchecked"
    assert "openpyxl is not installed" in outcome.detail
    assert cf.artifact().read_bytes() == body


def test_a_source_already_on_disk_has_its_text_checked_too(pinned, served, monkeypatch):
    on_disk(pinned, BODY)
    assert F.fetch_one(pinned).state == "present"
    changed(monkeypatch)
    assert F.fetch_one(pinned).state == "text differs"


def test_the_three_failures_are_three_outcomes_and_each_fails_the_command(
    pinned, served, monkeypatch, capsys
):
    claims = str(pinned.path.parent)
    served[URL] = LANDING
    assert F.main(["--claims", claims]) == 1
    assert "differs" in capsys.readouterr().out

    served[URL] = BODY
    changed(monkeypatch)
    assert F.main(["--claims", claims]) == 1
    out = capsys.readouterr().out
    assert "text differs" in out and "match their pin" in out
    assert "still not readable here" not in out, "the file is here; the reading is in question"

    monkeypatch.undo()
    assert F.fetch_one(pinned).state == "present"


def test_a_source_recording_no_derived_text_is_fetched_as_it_always_was(tmp_path, served):
    cf = claims_for(tmp_path, BODY, url=URL, extractor="jats-text")
    served[URL] = BODY
    assert F.fetch_one(cf).state == "fetched"


# --- pin: the digest nobody computes by hand ------------------------------------------------------


def test_the_first_quotation_pinned_records_the_version_and_the_text_it_was_found_in(tmp_path):
    cf = claims_for(tmp_path, BODY, extractor="jats-text")
    source = on_disk(cf, BODY)
    code = pin.main([str(cf.path), "--id", "design", "--quote", SENTENCE])
    written = load_claim_file(cf.path)
    assert code == 0
    assert written.source.extractor_version == 1
    assert written.source.derived_sha256 == derived_of(source, "jats-text")
    assert written.claims["design"].quotes[0].text == SENTENCE


def test_a_quotation_is_refused_once_the_extractor_produces_other_text(tmp_path, monkeypatch):
    # Written, the file would hold quotations taken from two readings under one digest.
    cf = claims_for(tmp_path, BODY, extractor="jats-text")
    on_disk(cf, BODY)
    assert pin.main([str(cf.path), "--id", "design", "--quote", SENTENCE]) == 0
    before = cf.path.read_text()
    changed(monkeypatch)
    V.clear_caches()
    with pytest.raises(pin.PinRefused, match="now produces other text"):
        pin.main([str(cf.path), "--id", "table", "--quote", "UK Biobank | 34,541"])
    assert cf.path.read_text() == before
