"""Built-in extractors for the formats a source arrives in that are neither a PDF nor text.

A claims file pins one file by sha256, and `fetch` installs a download only when its bytes are
the pinned ones. That works where the pinned file is the publisher's file. It fails where the
pinned file is a text somebody's script made from one -- a supplementary workbook rendered to
rows, a JATS article with its tags stripped -- because a reader's download is the workbook and
never the text. Measured on one audit's 199 claims files, on an empty checkout: 2 fetched, 156
differ, 41 unavailable. 2,688 of its 4,499 quotations are rows of supplementary spreadsheets.

The remedy is to pin the original and name the extraction, which is what a PDF already gets:
the pin names the publisher's bytes and `pdftotext` is recorded as what read them. These are
the same thing for four more formats.

    sheet-rows   .xlsx, .xls, .csv   one row per line, cells joined with ` | `
    docx-text    .docx               the body in document order, then notes, then headers
                                     and footers; tracked changes accepted (version 2)
    jats-text    JATS XML            one block per line, inline markup dropped
    html-text    HTML                one block per line, tags dropped, entities decoded,
                                     `<script>` and `<style>` left out

Each has a name a claims file can write and a version. The version is of the output, not of
the code: it changes when the same bytes would produce different text, because a quotation
pinned against one rendering says nothing about another. A claims file naming a version this
build does not ship is `unchecked`, with the reason.

They run in this process and execute nothing, so they are outside the `extract_cmd` allowlist:
a claims file naming one names a reader this package ships, not a program on the machine.

`openpyxl`, `xlrd` and `python-docx` (which only `docx-text` version 1 needs) are optional, as the PDF readers are. A missing one is
reported and never raised as an `ImportError` out of a check; the quotation is `unchecked`,
which is the outcome for a source nothing here could read.

What every extractor shares: whitespace inside a cell or a block is collapsed, so one row or
one paragraph is one line; an empty cell is dropped unless the source asks to keep them; a row
or block with no text produces no line. Nothing is reordered and nothing is summarized.
"""

from __future__ import annotations

import csv
import dataclasses
import datetime
import html.parser
import pathlib
import posixpath
import shlex
import warnings
import xml.etree.ElementTree as ET
import zipfile
from collections.abc import Callable, Iterable

from citations.exceptions import SourceUnreadableError

try:  # optional: `pip install "citations[sheets]"`
    import openpyxl
except ImportError:  # pragma: no cover - exercised by the extras matrix, not the suite
    openpyxl = None

try:  # optional: `pip install "citations[sheets]"`
    import xlrd
except ImportError:  # pragma: no cover
    xlrd = None

try:  # optional: `pip install "citations[docx]"`
    import docx
except ImportError:  # pragma: no cover
    docx = None

#: Between two cells of one row. Spaced, so a cell's last word and the next cell's first never
#: weld into a word that is in neither, and visible, so a quotation spanning two cells shows
#: that it does.
CELL_SEPARATOR = " | "


def _line(text: str) -> str:
    """One line: a cell or a block with every run of whitespace, newlines included, as a space.

    A cell holding a line break would otherwise put half a row on the next line, where it reads
    as a row of its own.
    """
    return " ".join(text.split())


def _cell(value: object) -> str:
    """A cell's value as the text a reader of the sheet sees.

    A spreadsheet stores `46` as the float 46.0 and `0.3` as 0.30000000000000004. Excel shows
    fifteen significant digits, which is what `.15g` keeps, so a quotation typed from the
    sheet matches what is rendered here. A date with no time of day is a date.
    """
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else format(value, ".15g")
    if isinstance(value, datetime.datetime):
        if value.time() == datetime.time():
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    return _line(str(value))


def _rows(rows: Iterable[Iterable[object]], keep_empty: bool) -> list[str]:
    """Rows as lines. An empty cell is dropped, or kept where its position is the content.

    Trailing empty cells are dropped either way: a sheet's used range is as wide as its widest
    row, and keeping them would end every shorter row in a run of separators.
    """
    lines = []
    for row in rows:
        cells = [_cell(value) for value in row]
        while cells and not cells[-1]:
            cells.pop()
        if not keep_empty:
            cells = [c for c in cells if c]
        if any(cells):
            lines.append(CELL_SEPARATOR.join(cells))
    return lines


def _named(path: pathlib.Path, sheet: str | None, names: list[str]) -> list[str]:
    """The sheets to read: the one named, or all of them in workbook order."""
    if sheet is None:
        return names
    if sheet not in names:
        raise SourceUnreadableError(
            path, f"no sheet named {sheet!r}; the workbook holds {', '.join(map(repr, names))}"
        )
    return [sheet]


def _xlsx(path: pathlib.Path, sheet: str | None, keep_empty: bool) -> list[str]:
    if openpyxl is None:
        raise SourceUnreadableError(
            path, 'openpyxl is not installed -- pip install "citations[sheets]"'
        )
    # `data_only`: the value a formula last computed, which is what the sheet shows, and never
    # the formula. `read_only`: a supplementary workbook can be tens of megabytes.
    #
    # openpyxl warns, once per sheet, about workbook features it does not model -- a data
    # validation extension, a conditional format. None of them is a cell's value, and on a
    # publisher's workbook the warning lands in the middle of a report about quotations.
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            lines: list[str] = []
            for name in _named(path, sheet, book.sheetnames):
                ws = book[name]
                # A workbook written by something other than Excel can record a used range of
                # one cell, and a read-only sheet trusts it: every row after the first is not
                # returned and nothing says so. Discarding the recorded range makes it read
                # what is there.
                ws.reset_dimensions()
                lines += _rows(ws.iter_rows(values_only=True), keep_empty)
            return lines
        finally:
            book.close()


def _xls(path: pathlib.Path, sheet: str | None, keep_empty: bool) -> list[str]:
    if xlrd is None:
        raise SourceUnreadableError(
            path, 'xlrd is not installed -- pip install "citations[sheets]"'
        )
    book = xlrd.open_workbook(str(path))

    def value(cell) -> object:
        # The old format has no integer or date type. A date is a float tagged as one, read
        # against the workbook's epoch; a boolean is 0 or 1 tagged as one.
        if cell.ctype == xlrd.XL_CELL_DATE:
            return xlrd.xldate_as_datetime(cell.value, book.datemode)
        if cell.ctype == xlrd.XL_CELL_BOOLEAN:
            return bool(cell.value)
        if cell.ctype == xlrd.XL_CELL_ERROR:
            return xlrd.error_text_from_code.get(cell.value, "")
        return cell.value

    lines: list[str] = []
    for name in _named(path, sheet, book.sheet_names()):
        ws = book.sheet_by_name(name)
        lines += _rows(([value(c) for c in ws.row(r)] for r in range(ws.nrows)), keep_empty)
    return lines


def _csv(path: pathlib.Path, sheet: str | None, keep_empty: bool) -> list[str]:
    if sheet is not None:
        raise SourceUnreadableError(path, f"a .csv has no sheets, and the source names {sheet!r}")
    # `utf-8-sig`: a file exported from Excel begins with a byte-order mark, which would
    # otherwise be the first character of the first cell.
    with path.open(newline="", encoding="utf-8-sig", errors="replace") as handle:
        return _rows(csv.reader(handle), keep_empty)


_SHEET_READERS: dict[str, Callable[[pathlib.Path, str | None, bool], list[str]]] = {
    ".xlsx": _xlsx,
    ".xlsm": _xlsx,
    ".xls": _xls,
    ".csv": _csv,
}


def _sheet_rows(path: pathlib.Path, sheet: str | None, keep_empty: bool) -> str:
    """A workbook's rows, one per line, sheets in workbook order."""
    read = _SHEET_READERS.get(path.suffix.lower())
    if read is None:
        raise SourceUnreadableError(
            path, f"sheet-rows reads {', '.join(_SHEET_READERS)}, and this is {path.suffix!r}"
        )
    return "\n".join(read(path, sheet, keep_empty))


def _docx_text(path: pathlib.Path, sheet: str | None, keep_empty: bool) -> str:
    """Body paragraphs in order, then each table one row per line.

    Tables come after the paragraphs and not where they sit among them, which is the order
    `python-docx` exposes and the order the texts this replaces were made in. Headers, footers,
    footnotes, text boxes and a table nested inside a cell are not read: a quotation taken from
    one of those is `not found`, and the remedy is a different extractor, not a wider match.
    """
    if docx is None:
        raise SourceUnreadableError(
            path, 'python-docx is not installed -- pip install "citations[docx]"'
        )
    document = docx.Document(str(path))
    lines = [_line(p.text) for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            # A merged cell is returned once per column it spans. Without this a row reads
            # `Cohort | Cohort | Cohort | 1,204`, and a quotation of the row as printed fails.
            cells, last = [], None
            for cell in row.cells:
                if cell._tc is not last:
                    cells.append(cell.text)
                last = cell._tc
            lines += _rows([cells], keep_empty)
    return "\n".join(line for line in lines if line)


# --- docx-text version 2: the document's own XML ---------------------------------------------

#: The most XML, uncompressed, read from one `.docx` across every part. Generous for a
#: manuscript or a supplement, whose body is a few megabytes; a part that inflates past it is
#: refused rather than read into memory, which is what a zip bomb relies on.
DOCX_MAX_XML_BYTES = 128 * 1024 * 1024

#: The namespaces a part may be written in, transitional and strict, by the prefix used below.
_DOCX_NS = {
    "http://schemas.openxmlformats.org/wordprocessingml/2006/main": "w",
    "http://purl.oclc.org/ooxml/wordprocessingml/main": "w",
    "http://schemas.openxmlformats.org/officeDocument/2006/math": "m",
    "http://purl.oclc.org/ooxml/officeDocument/math": "m",
    "http://schemas.openxmlformats.org/markup-compatibility/2006": "mc",
}
_DOCX_RELS_NS = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "http://purl.oclc.org/ooxml/officeDocument/relationships",
)

#: Inside a paragraph: elements whose children are read as if they stood in its place. An
#: insertion and a move's destination are read, because the text is the accepted view.
_DOCX_TRANSPARENT = frozenset(
    {
        "w:hyperlink",
        "w:ins",
        "w:moveTo",
        "w:smartTag",
        "w:customXml",
        "w:fldSimple",
        "w:sdt",
        "w:sdtContent",
        "w:dir",
        "w:bdo",
    }
)

#: Never read: a deletion and a move's origin, which the accepted view removes, and the
#: properties of a content control, which hold its settings and not its text.
_DOCX_REMOVED = frozenset({"w:del", "w:moveFrom", "w:sdtPr", "w:sdtEndPr"})

#: What a note that is not one of the document's notes is: the rule and the continuation rule
#: Word draws above the notes on a page.
_DOCX_NOTE_FURNITURE = frozenset({"separator", "continuationSeparator", "continuationNotice"})


def _docx_key(element: ET.Element) -> str:
    """`w:p`, `m:t` or `mc:Choice` for an element in a namespace read here, else empty."""
    namespace, _, local = element.tag.rpartition("}")
    prefix = _DOCX_NS.get(namespace.lstrip("{"))
    return f"{prefix}:{local}" if prefix else ""


def _docx_attr(element: ET.Element, name: str) -> str | None:
    """A `w:` attribute, in whichever namespace the part is written."""
    for namespace, prefix in _DOCX_NS.items():
        if prefix == "w" and (value := element.get(f"{{{namespace}}}{name}")) is not None:
            return value
    return None


def _docx_rid(element: ET.Element) -> str | None:
    """The relationship an element points at with `r:id`."""
    for namespace in _DOCX_RELS_NS:
        if (value := element.get(f"{{{namespace}}}id")) is not None:
            return value
    return None


class _DocxPackage:
    """The parts of a `.docx`, each read once, under one budget of uncompressed bytes."""

    def __init__(self, path: pathlib.Path, archive: zipfile.ZipFile) -> None:
        self.path = path
        self.archive = archive
        self.left = DOCX_MAX_XML_BYTES

    def xml(self, name: str) -> ET.Element | None:
        """The root of a part, or None where the package holds no part of that name.

        A part that declares a document type is refused: no Office part carries one, and a
        declaration is where entity expansion lives. A part that is not UTF-8 is refused for
        the same reason, since the check for a declaration reads the text.
        """
        try:
            info = self.archive.getinfo(name)
        except KeyError:
            return None
        with self.archive.open(info) as handle:
            data = handle.read(self.left + 1)
        if len(data) > self.left:
            raise SourceUnreadableError(
                self.path,
                f"its XML inflates past {DOCX_MAX_XML_BYTES:,} bytes; it is refused unread",
            )
        self.left -= len(data)
        try:
            text = data.decode("utf-8")
        except UnicodeDecodeError as e:
            raise SourceUnreadableError(self.path, f"{name} is not UTF-8") from e
        if "<!doctype" in text.lower() or "<!entity" in text.lower():
            raise SourceUnreadableError(
                self.path, f"{name} declares a document type, which no Word part does"
            )
        try:
            return ET.fromstring(text)
        except ET.ParseError as e:
            raise SourceUnreadableError(self.path, f"{name} is not well-formed XML: {e}") from e

    def related(self, part: str) -> dict[str, tuple[str, str]]:
        """A part's relationships: id to (type, the part it names inside the package).

        A target outside the package, `TargetMode="External"`, is left out: nothing here is
        fetched.
        """
        folder, _, file = part.rpartition("/")
        root = self.xml(f"{folder}/_rels/{file}.rels" if folder else f"_rels/{file}.rels")
        found: dict[str, tuple[str, str]] = {}
        for rel in root if root is not None else ():
            target, kind = rel.get("Target", ""), rel.get("Type", "")
            if rel.get("TargetMode") == "External" or not target:
                continue
            inside = target.lstrip("/") if target.startswith("/") else f"{folder}/{target}"
            found[rel.get("Id", "")] = (kind, posixpath.normpath(inside).lstrip("/"))
        return found


class _DocxText:
    """The walk over a part's XML that writes its text, one block per line."""

    def __init__(self) -> None:
        self.fields: list[bool] = []
        """One entry per complex field open: True while its instruction is being read, False
        once its result is. Text is written only where no instruction is open."""
        self.notes: dict[str, list[str]] = {"footnote": [], "endnote": []}
        """The id of each note the text refers to, in the order it first refers to it."""
        self.carry = ""
        """The text of a paragraph whose mark a tracked change deletes, which joins the next."""

    @property
    def shown(self) -> bool:
        return not any(self.fields)

    def blocks(self, container: ET.Element, out: list[str], tables: list[str] | None = None):
        """Paragraphs and tables, in the order they stand. A table inside a table cell goes to
        `tables`, where one is given, so its rows are not welded into the cell."""
        for child in container:
            key = _docx_key(child)
            if key == "w:p":
                self.paragraph(child, out)
            elif key == "w:tbl":
                self.table(child, out if tables is None else tables)
            elif key in ("w:sdt", "w:sdtContent", "w:customXml", "w:ins", "w:moveTo"):
                self.blocks(child, out, tables)
            elif key == "mc:AlternateContent" and (chosen := self.choice(child)) is not None:
                self.blocks(chosen, out, tables)
        if self.carry:
            out.append(_line(self.carry))
            self.carry = ""

    def paragraph(self, p: ET.Element, out: list[str]) -> None:
        carry, self.carry = self.carry, ""
        parts: list[str] = []
        boxes: list[str] = []
        self.inline(p, parts, boxes)
        text = carry + "".join(parts)
        mark = next((c for c in p if _docx_key(c) == "w:pPr"), None)
        deleted = mark is not None and any(
            _docx_key(c) == "w:del" for r in mark if _docx_key(r) == "w:rPr" for c in r
        )
        if deleted:
            self.carry = text
        else:
            out.append(_line(text))
        out.extend(boxes)

    def inline(self, element: ET.Element, parts: list[str], boxes: list[str]) -> None:
        for child in element:
            key = _docx_key(child)
            if key == "w:r":
                self.run(child, parts, boxes)
            elif key in _DOCX_REMOVED:
                continue
            elif key in _DOCX_TRANSPARENT:
                self.inline(child, parts, boxes)
            elif key in ("m:oMath", "m:oMathPara"):
                if self.shown:
                    parts.append(
                        "".join(t.text or "" for t in child.iter() if _docx_key(t) == "m:t")
                    )
            elif key == "mc:AlternateContent" and (chosen := self.choice(child)) is not None:
                self.inline(chosen, parts, boxes)

    def run(self, r: ET.Element, parts: list[str], boxes: list[str]) -> None:
        for child in r:
            key = _docx_key(child)
            if key == "w:fldChar":
                kind = _docx_attr(child, "fldCharType")
                if kind == "begin":
                    self.fields.append(True)
                elif kind == "separate" and self.fields:
                    self.fields[-1] = False
                elif kind == "end" and self.fields:
                    self.fields.pop()
            elif not self.shown:
                continue
            elif key == "w:t":
                parts.append(child.text or "")
            elif key in ("w:tab", "w:ptab", "w:br", "w:cr"):
                parts.append(" ")
            elif key == "w:noBreakHyphen":
                parts.append("-")
            elif key in ("w:footnoteReference", "w:endnoteReference"):
                ids = self.notes[key[2:-9]]
                if (note := _docx_attr(child, "id")) is not None and note not in ids:
                    ids.append(note)
            elif key in ("m:oMath", "m:oMathPara"):
                parts.append("".join(t.text or "" for t in child.iter() if _docx_key(t) == "m:t"))
            elif key == "mc:AlternateContent":
                if (chosen := self.choice(child)) is not None:
                    self.text_boxes(chosen, boxes)
            elif key not in ("w:rPr", "w:instrText", "w:delText", "w:delInstrText"):
                # A drawing, a VML shape or an embedded object: the text boxes inside it.
                self.text_boxes(child, boxes)

    def text_boxes(self, element: ET.Element, out: list[str]) -> None:
        for child in element:
            key = _docx_key(child)
            if key == "w:txbxContent":
                self.blocks(child, out)
            elif key == "mc:AlternateContent":
                if (chosen := self.choice(child)) is not None:
                    self.text_boxes(chosen, out)
            else:
                self.text_boxes(child, out)

    def table(self, tbl: ET.Element, out: list[str]) -> None:
        nested: list[str] = []
        for row in self._rows_of(tbl):
            properties = next((c for c in row if _docx_key(c) == "w:trPr"), None)
            if properties is not None and any(_docx_key(c) == "w:del" for c in properties):
                continue  # a row the accepted view deletes
            cells = []
            for cell in self._cells_of(row):
                lines: list[str] = []
                self.blocks(cell, lines, nested)
                cells.append(" ".join(line for line in lines if line))
            out.extend(_rows([cells], keep_empty=False))
            out.extend(nested)
            nested.clear()

    def _rows_of(self, element: ET.Element) -> Iterable[ET.Element]:
        for child in element:
            key = _docx_key(child)
            if key == "w:tr":
                yield child
            elif key in ("w:sdt", "w:sdtContent", "w:customXml", "w:ins", "w:moveTo"):
                yield from self._rows_of(child)

    def _cells_of(self, row: ET.Element) -> Iterable[ET.Element]:
        for child in row:
            key = _docx_key(child)
            if key == "w:tc":
                yield child
            elif key in ("w:sdt", "w:sdtContent", "w:customXml", "w:ins", "w:moveTo"):
                yield from self._cells_of(child)

    @staticmethod
    def choice(alternate: ET.Element) -> ET.Element | None:
        """The one branch of `mc:AlternateContent` to read: the first `mc:Choice`, else the
        `mc:Fallback`. The branches are two encodings of one object, and reading both would
        print a text box twice."""
        branches = [c for c in alternate if _docx_key(c) in ("mc:Choice", "mc:Fallback")]
        chosen = [c for c in branches if _docx_key(c) == "mc:Choice"] or branches
        return chosen[0] if chosen else None


def _docx_xml_text(path: pathlib.Path, sheet: str | None, keep_empty: bool) -> str:
    """`docx-text` version 2: a `.docx` read from its own XML, in this order.

      1. The body, in document order: each paragraph one line, and each table where it stands,
         one row per line, cells left to right joined with ` | `, a merged cell once. A table
         inside a cell follows the row holding it. A text box follows the paragraph its shape
         is anchored in. A content control's content stands where the control does.
      2. Each footnote, then each endnote, in the order the text first refers to it, numbered
         from 1 in that order and marked `Footnote 1.` or `Endnote 1.`. That is Word's default
         numbering; a custom start, format or mark is not rendered. The reference mark in the
         body is not written there.
      3. Each header and footer part once, in the order the sections name them.

    Within a paragraph: the runs, a hyperlink's text, a content control's text, a simple
    field's result and a complex field's result (the text between its `separate` and its
    `end`; the instruction is left out), and an equation's characters in order (OMML `m:t`, with
    no fraction, script or radical rendered). Tracked changes are read as the accepted view:
    an insertion and a move's destination are read, a deletion and a move's origin are not, a
    row a change deletes is left out, and a paragraph whose mark a change deletes joins the
    next. A tab or a line break is a space, and whitespace is collapsed as every extractor here
    collapses it.

    Not read: comments, which are a reviewer's text and not the document's; automatic list and
    heading numbers, which Word computes and does not store; a footnote nothing refers to; the
    glossary; an embedded document (`w:altChunk`) or object's own contents; and anything an
    image shows. Hidden text (`w:vanish`) is read.

    Parsed with the standard library from the zip, never through a file path inside it and
    with nothing fetched. A part that declares a document type, or is not UTF-8, is refused,
    and so is a package whose XML inflates past `DOCX_MAX_XML_BYTES`.
    """
    try:
        archive = zipfile.ZipFile(path)
    except zipfile.BadZipFile as e:
        raise SourceUnreadableError(path, "not a .docx: it is not a zip archive") from e
    with archive:
        package = _DocxPackage(path, archive)
        main = next(
            (
                target
                for kind, target in package.related("").values()
                if kind.endswith("/officeDocument")
            ),
            "word/document.xml",
        )
        document = package.xml(main)
        if document is None:
            raise SourceUnreadableError(path, f"not a .docx: it holds no {main}")
        related = package.related(main)
        walk = _DocxText()
        lines: list[str] = []
        body = next((c for c in document if _docx_key(c) == "w:body"), None)
        if body is not None:
            walk.blocks(body, lines)

        for kind, label in (("footnote", "Footnote"), ("endnote", "Endnote")):
            part = next((t for k, t in related.values() if k.endswith(f"/{kind}s")), None)
            root = package.xml(part) if part else None
            if root is None:
                continue
            notes = {
                _docx_attr(note, "id"): note
                for note in root
                if _docx_key(note) == f"w:{kind}"
                and _docx_attr(note, "type") not in _DOCX_NOTE_FURNITURE
            }
            for number, note_id in enumerate(walk.notes[kind], start=1):
                if (note := notes.get(note_id)) is None:
                    continue
                walk.fields = []
                text: list[str] = []
                walk.blocks(note, text)
                text = [line for line in text if line]
                if text:
                    lines.append(f"{label} {number}. {text[0]}")
                    lines += text[1:]

        # Each header and footer once, in the order the sections name them.
        seen: list[str] = []
        for reference in document.iter():
            if _docx_key(reference) in ("w:headerReference", "w:footerReference"):
                target = related.get(_docx_rid(reference) or "", ("", ""))[1]
                if target and target not in seen:
                    seen.append(target)
        for target in seen:
            if (root := package.xml(target)) is not None:
                walk.fields = []
                walk.blocks(root, lines)
    return "\n".join(line for line in lines if line)


#: JATS elements that sit inside a sentence. Everything else starts a new line. Listed this way
#: round on purpose: an inline element missing from here splits a sentence across two lines,
#: which the matcher's whitespace folding absorbs, while a block element missing from a list of
#: blocks would have its text dropped, which reads as a passage the source does not contain.
_JATS_INLINE = frozenset(
    {
        "abbrev",
        "bold",
        "email",
        "ext-link",
        "inline-formula",
        "inline-graphic",
        "italic",
        "monospace",
        "named-content",
        "overline",
        "roman",
        "sans-serif",
        "sc",
        "strike",
        "styled-content",
        "sub",
        "sup",
        "underline",
        "uri",
        "xref",
    }
)


def _tag(element: ET.Element) -> str:
    """The element's name without a namespace, which MathML and some deposits carry."""
    return element.tag.rsplit("}", 1)[-1]


def _jats_lines(element: ET.Element, lines: list[str]) -> None:
    """Append the text under `element`, one block per line, in document order."""
    if _tag(element) == "tr":
        cells = ["".join(c.itertext()) for c in element if _tag(c) in ("td", "th")]
        lines += _rows([cells], keep_empty=False)
        return
    run = [element.text or ""]
    for child in element:
        if _tag(child) == "break":
            run.append(" ")
        elif _tag(child) in _JATS_INLINE:
            run.append("".join(child.itertext()))
        else:
            lines.append(_line("".join(run)))
            run = []
            _jats_lines(child, lines)
        run.append(child.tail or "")
    lines.append(_line("".join(run)))


def _jats_text(path: pathlib.Path, sheet: str | None, keep_empty: bool) -> str:
    """A JATS article as text: one block per line, inline markup dropped, entities decoded.

    Read as plain text, the XML keeps its markup, and a quotation crossing `<italic>`, `<sup>`,
    `<xref>` or a character entity does not resolve. Measured on 72 Europe PMC and NCBI
    articles: 194 of 314 quotations resolved in the XML as served, and 308 resolve in this
    text. The other 6 are table rows quoted with spaces between their cells.

    The whole document is read, front matter and reference list included: leaving a part out
    is a decision about what a paper may be quoted on, and that is the claims file's to make.
    A table is one row per line with cells joined by ` | `; spans are not expanded and header
    cells are not repeated beside the body cells.
    """
    try:
        root = ET.parse(path).getroot()
    except ET.ParseError as e:
        raise SourceUnreadableError(path, f"not well-formed XML: {e}") from e
    lines: list[str] = []
    _jats_lines(root, lines)
    return "\n".join(line for line in lines if line)


#: HTML elements that sit inside a sentence. Everything else starts a new line, for the reason
#: `_JATS_INLINE` gives: an inline element missing from here splits a sentence across two lines,
#: which the matcher's whitespace folding absorbs, and no element's text is ever dropped for
#: being unlisted.
_HTML_INLINE = frozenset(
    {
        "a",
        "abbr",
        "b",
        "bdi",
        "bdo",
        "cite",
        "code",
        "data",
        "del",
        "dfn",
        "em",
        "font",
        "i",
        "ins",
        "kbd",
        "label",
        "mark",
        "q",
        "s",
        "samp",
        "small",
        "span",
        "strong",
        "sub",
        "sup",
        "time",
        "tt",
        "u",
        "var",
        "wbr",
    }
)

#: Elements whose text is a program or a stylesheet and no part of what the page says.
_HTML_UNREAD = frozenset({"script", "style"})

_HTML_CELLS = frozenset({"td", "th"})


class _HTMLText(html.parser.HTMLParser):
    """The text of an HTML page, gathered a block at a time.

    Character references are decoded by the parser before the text arrives, and a comment
    never arrives at all.
    """

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.lines: list[str] = []
        self._run: list[str] = []
        self._cells: list[str] | None = None
        self._unread = 0

    def _flush(self) -> None:
        text, self._run = _line("".join(self._run)), []
        if self._cells is not None:
            self._cells.append(text)
        elif text:
            self.lines.append(text)

    def _end_row(self) -> None:
        self._flush()
        if self._cells is not None:
            self.lines += _rows([self._cells], keep_empty=False)
        self._cells = None

    def _boundary(self, tag: str, opening: bool) -> None:
        if tag in _HTML_UNREAD:
            self._unread = max(0, self._unread + (1 if opening else -1))
        elif tag == "tr":
            self._end_row()
            if opening:
                self._cells = []
        elif tag in _HTML_CELLS or self._cells is None:
            self._flush()
        else:
            # A block inside a table cell: its text stays in the cell, parted by a space.
            self._run.append(" ")

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag not in _HTML_INLINE:
            self._boundary(tag, opening=True)

    def handle_endtag(self, tag: str) -> None:
        if tag not in _HTML_INLINE:
            self._boundary(tag, opening=False)

    def handle_data(self, data: str) -> None:
        if not self._unread:
            self._run.append(data)

    def close(self) -> None:
        super().close()
        self._end_row()


def _html_text(path: pathlib.Path, sheet: str | None, keep_empty: bool) -> str:
    """An HTML page as the text it shows: one block per line, tags dropped, entities decoded.

    Read as plain text, a page keeps its markup, and a quotation crossing `<a>` or `<em>`, or
    holding a character the page writes as `&amp;` or `&rsquo;`, does not resolve. On a web
    page that is most sentences.

    An inline element is dropped and its text kept in place, so a sentence reads through a
    link. Every other element ends a line, so the last word of one block and the first of the
    next are never welded into a word neither holds. A line break is white space to the
    matcher, as it is in any text source, so a quotation running from the end of one block
    into the start of the next resolves as it does across a line break in a `.txt`. A table is
    one row per line with cells joined by ` | `, as the other extractors write one, so a
    quotation spanning two cells resolves only where it shows that it does.

    The text of `<script>` and `<style>` is left out. Nothing else is: navigation, footers and
    text a stylesheet hides are in the page, and leaving a part out is a decision about what a
    page may be quoted on. Read as UTF-8, with a byte that is not replaced.
    """
    parser = _HTMLText()
    parser.feed(path.read_text(encoding="utf-8", errors="replace"))
    parser.close()
    return "\n".join(parser.lines)


@dataclasses.dataclass(frozen=True)
class Extractor:
    """One built-in extractor: the name a claims file writes, and the version of its output."""

    name: str
    version: int
    read: Callable[[pathlib.Path, str | None, bool], str]

    takes_options: bool = False
    """Whether `--sheet` and `--keep-empty` mean anything to it. Only a workbook has sheets,
    and an option given to an extractor that ignores it is a declaration nothing honoured."""


EXTRACTORS: dict[str, Extractor] = {
    e.name: e
    for e in (
        Extractor("sheet-rows", 1, _sheet_rows, takes_options=True),
        Extractor("docx-text", 2, _docx_xml_text),
        Extractor("jats-text", 1, _jats_text),
        Extractor("html-text", 1, _html_text),
    )
}

#: Earlier versions of an extractor that a claims file may still name, by name and version. A
#: source pinned under one goes on being read by it, so its `derived_sha256` still describes the
#: text it gets: a newer rendering is a different text, and quotations pinned against the older
#: one say nothing about it. A source that names no version, and one read by its suffix, gets the
#: version in `EXTRACTORS`; `citations pin` records that version on the first quotation.
EARLIER: dict[tuple[str, str], Extractor] = {
    (e.name, str(e.version)): e
    for e in (
        # Body paragraphs, then tables, through `python-docx`: no notes, headers, text boxes,
        # content controls, field results or tracked insertions. Kept for the sources pinned
        # under it.
        Extractor("docx-text", 1, _docx_text),
    )
}

#: The extractor a format obviously wants, chosen where a source declares nothing. The same
#: rule `verify.BY_SUFFIX` applies to `.tex`: without it a workbook reached the PDF readers and
#: every quotation in it was `unchecked`. `.xml` is not here. An XML source that declares
#: nothing is read as plain text, as it always was, and quotations already pinned against the
#: markup go on resolving; `jats-text` is for a source that names it. `.html` is not here for
#: the same reason: `html-text` is for a source that names it, which `citations pin --url`
#: does for a page it pins.
BY_SUFFIX: dict[str, str] = {
    ".xlsx": "sheet-rows",
    ".xlsm": "sheet-rows",
    ".xls": "sheet-rows",
    ".csv": "sheet-rows",
    ".docx": "docx-text",
}


@dataclasses.dataclass(frozen=True)
class Declared:
    """A built-in extractor as one source declared it, options included."""

    extractor: Extractor
    sheet: str | None = None
    keep_empty: bool = False

    @property
    def name(self) -> str:
        """What a result records: the extractor, its version, and the options that shaped it.

        Always with the version, whether or not the claims file wrote one. Two versions of one
        extractor are two readings of the same bytes, and a result that does not say which
        cannot be compared with a later one.
        """
        words = [f"{self.extractor.name}@{self.extractor.version}"]
        if self.sheet is not None:
            words += ["--sheet", self.sheet]
        if self.keep_empty:
            words.append("--keep-empty")
        return shlex.join(words)

    def read(self, path: pathlib.Path) -> str:
        """The text this extractor produces from `path`, or `SourceUnreadableError` saying why.

        Any failure to parse is the document's and is reported as a source that could not be
        read. It is never an empty text, which would say the document holds no words.
        """
        try:
            return self.extractor.read(path, self.sheet, self.keep_empty)
        except SourceUnreadableError:
            raise
        except Exception as e:
            raise SourceUnreadableError(
                path, f"{self.extractor.name} could not read it: {e}"
            ) from e


def is_builtin(extractor: str) -> bool:
    """Whether a declared reading, or the name recorded on a result, is one of `EXTRACTORS`."""
    words = extractor.split()
    return bool(words) and words[0].partition("@")[0] in EXTRACTORS


def declared(path: pathlib.Path, text: str) -> Declared | None:
    """The built-in extractor `text` names, or None where it names a program instead.

    `text` is `name`, or `name@version`, then `--sheet NAME` and `--keep-empty` where the
    extractor takes them: the form `ClaimSource.reader` writes from a source's `extractor`,
    `extractor_version`, `sheet` and `empty_cells` fields.

    Raises `SourceUnreadableError`, so the quotation is `unchecked`, where the version is not
    the one shipped here or an option is not one the extractor takes. Reading with a different
    version, or ignoring the option, would produce a text the claims file did not ask for and
    report a verdict against it.
    """
    try:
        words = shlex.split(text)
    except ValueError:
        return None  # not ours to report: `verify._argv` says why it will not parse
    if not words:
        return None
    name, _, version = words[0].partition("@")
    extractor = EXTRACTORS.get(name)
    if extractor is None:
        return None
    if version and version != str(extractor.version):
        earlier = EARLIER.get((name, version))
        if earlier is None:
            shipped = sorted(
                [extractor.version, *(e.version for k, e in EARLIER.items() if k[0] == name)]
            )
            raise SourceUnreadableError(
                path,
                f"the source names {name} version {version} and this build ships version "
                f"{' and '.join(map(str, shipped))}. The two produce different text from the same "
                f"bytes, so nothing was read: install the citations release that ships version "
                f"{version}, or pin the quotations again under version {extractor.version}",
            )
        extractor = earlier
    sheet, keep_empty, rest = None, False, words[1:]
    while rest:
        word = rest.pop(0)
        if extractor.takes_options and word == "--sheet" and rest:
            sheet = rest.pop(0)
        elif extractor.takes_options and word == "--keep-empty":
            keep_empty = True
        else:
            raise SourceUnreadableError(path, f"{name} does not take {word!r}")
    return Declared(extractor, sheet, keep_empty)


__all__ = [
    "BY_SUFFIX",
    "CELL_SEPARATOR",
    "DOCX_MAX_XML_BYTES",
    "EARLIER",
    "EXTRACTORS",
    "Declared",
    "Extractor",
    "declared",
    "is_builtin",
]
