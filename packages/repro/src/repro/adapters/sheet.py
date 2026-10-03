"""Excel workbooks, addressed by sheet, column and a row predicate.

`.xlsx` and `.xlsm` are read with `openpyxl` and `.xls` with `xlrd`, both optional
(`pip install "reproducible-science[sheets]"`). Without them a check is `unchecked` and says
which extra to install; it never fails and never passes.

A cell is read as the value a formula last computed, which is what the sheet shows, and never
as the formula. A workbook saved by a program that does not compute formulas holds no computed
value for them, and such a cell is reported as holding no value: a check that substituted the
formula text, or zero, would be comparing the manuscript against something the file does not
say.
"""

from __future__ import annotations

import pathlib
import warnings
import zipfile

from repro.adapters.base import Found, Resolution, _no
from repro.adapters.cells import Cell, render
from repro.adapters.rows import resolve_rows
from repro.exceptions import ArtifactUnreadableError, BackendUnavailableError
from repro.models import SheetLocator

try:  # optional: `pip install "reproducible-science[sheets]"`
    import openpyxl
except ImportError:  # pragma: no cover - exercised by the extras matrix, not the suite
    openpyxl = None

try:  # optional: `pip install "reproducible-science[sheets]"`
    import xlrd
    import xlrd.biffh
except ImportError:  # pragma: no cover
    xlrd = None

WORKBOOK_SUFFIXES = {".xlsx", ".xlsm", ".xls"}

_INSTALL = 'pip install "reproducible-science[sheets]"'


def _xlsx(path: pathlib.Path, sheet: str | None) -> tuple[list[str], list[list[object]]]:
    """Sheet names, and the rows of the one selected; `sheet` is resolved by `_choose`."""
    if openpyxl is None:
        raise BackendUnavailableError("sheet", f"openpyxl is not installed -- {_INSTALL}")
    # openpyxl warns about workbook features it does not model -- a data validation extension,
    # a conditional format. None is a cell's value, and the warning would land in the middle of
    # a verification report.
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        try:
            book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        except (zipfile.BadZipFile, KeyError, OSError, ValueError) as e:
            raise ArtifactUnreadableError(path, f"not a readable workbook: {e}") from e
        try:
            names = list(book.sheetnames)
            chosen = _choose(names, sheet)
            if chosen is None:
                return names, []
            ws = book[chosen]
            # A workbook written by something other than Excel can record a used range of one
            # cell, and a read-only sheet trusts it: every row after the first is silently not
            # returned. The citations extractor met this on publishers' supplements.
            ws.reset_dimensions()
            return names, [list(row) for row in ws.iter_rows(values_only=True)]
        finally:
            book.close()


def _xls(path: pathlib.Path, sheet: str | None) -> tuple[list[str], list[list[object]]]:
    if xlrd is None:
        raise BackendUnavailableError("sheet", f"xlrd is not installed -- {_INSTALL}")
    try:
        book = xlrd.open_workbook(str(path))
    except (xlrd.biffh.XLRDError, OSError, EOFError, AssertionError) as e:
        raise ArtifactUnreadableError(path, f"not a readable workbook: {e}") from e
    names = book.sheet_names()
    chosen = _choose(names, sheet)
    if chosen is None:
        return names, []
    ws = book.sheet_by_name(chosen)

    def value(cell) -> object:
        # The old format has no integer type, and an empty cell is the empty string. A date is
        # a float tagged as one; it is not a number a manuscript reports, and reading it as
        # one would compare a day count against a printed value.
        if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
            return None
        if cell.ctype == xlrd.XL_CELL_DATE:
            return xlrd.xldate_as_datetime(cell.value, book.datemode)
        if cell.ctype == xlrd.XL_CELL_BOOLEAN:
            return bool(cell.value)
        if cell.ctype == xlrd.XL_CELL_ERROR:
            return xlrd.error_text_from_code.get(cell.value, "#ERROR")
        return cell.value

    return names, [[value(c) for c in ws.row(r)] for r in range(ws.nrows)]


def _choose(names: list[str], sheet: str | None) -> str | None:
    """The sheet a locator selects, or None when it selects none; `_resolve_sheet` says why."""
    if sheet is None:
        return names[0] if len(names) == 1 else None
    return sheet if sheet in names else None


def _frame(raw: list[list[object]]) -> tuple[list[str], list[dict[str, Cell]]]:
    """The first row as the header and every later row keyed by it, as delimited text is read.

    A column whose header cell is empty is left out: no locator can name it, and a sheet's used
    range routinely runs past its last named column. Kept as columns named "", two of them
    would repeat a header name and the repeated-header rule would refuse the whole sheet.
    """
    if not raw:
        return [], []
    names = [h if isinstance(h, str) else "" for h in map(render, raw[0])]
    named = [(i, n) for i, n in enumerate(names) if n]
    header = [n for _, n in named]
    rows: list[dict[str, Cell]] = [
        {n: render(values[i]) if i < len(values) else None for i, n in named} for values in raw[1:]
    ]
    return header, rows


def _resolve_sheet(locator: SheetLocator, path: pathlib.Path) -> Found:
    suffix = path.suffix.lower()
    if suffix not in WORKBOOK_SUFFIXES:
        return _no(
            Resolution.FORMAT_UNSUPPORTED,
            f"a sheet locator addresses an Excel workbook; {path.name} is "
            f"{path.suffix or 'extensionless'}",
        )
    read = _xls if suffix == ".xls" else _xlsx
    names, raw = read(path, locator.sheet)
    if locator.sheet is None and len(names) != 1:
        return _no(
            Resolution.SELECTOR_INVALID,
            f"{path.name} holds {len(names)} sheets; name one of {', '.join(names[:8])}",
        )
    if locator.sheet is not None and locator.sheet not in names:
        return _no(
            Resolution.SELECTOR_INVALID,
            f"{path.name} has no sheet {locator.sheet!r}; sheets are {', '.join(names[:8])}",
        )
    header, rows = _frame(raw)
    name = f"{path.name}, sheet {locator.sheet or names[0]!r}"
    wanted = {k: render(v) for k, v in locator.where.items()}
    return resolve_rows(header, rows, locator.column, wanted, name)
